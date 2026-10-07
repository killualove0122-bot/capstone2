import pptx
import json

prs = pptx.Presentation(r'd:\캡스톤판넬 1차_샘플.pptx')
slide = prs.slides[0]

def dump_shape(shape):
    info = {
        'name': shape.name,
        'type': str(shape.shape_type),
        'left': shape.left.inches if hasattr(shape, 'left') and shape.left is not None else 0,
        'top': shape.top.inches if hasattr(shape, 'top') and shape.top is not None else 0,
        'width': shape.width.inches if hasattr(shape, 'width') and shape.width is not None else 0,
        'height': shape.height.inches if hasattr(shape, 'height') and shape.height is not None else 0,
    }
    if shape.has_text_frame:
        info['text'] = shape.text_frame.text
    if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        info['subshapes'] = [dump_shape(s) for s in shape.shapes]
    return info

data = [dump_shape(s) for s in slide.shapes]
with open('sample_structure.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print(f"Dumped {len(data)} shapes to sample_structure.json")
